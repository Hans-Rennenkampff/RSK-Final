from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles as SF
import uvicorn
from fastapi.middleware.cors import CORSMiddleware
import pandas as pd
from openpyxl import load_workbook
import matplotlib.pyplot as plt
import numpy as np
import  os
app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Измените на конкретный URL в продакшене
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


app.mount("/static", SF(directory="c:/RSK-master/resourses/web"), name="static")

@app.get("/theme/{theme_name}.css")
async def get_theme_css(theme_name: str):
    path = f"resourses/web/theme-{theme_name}.css"
    return FileResponse(path)

@app.get("/")
async def read_index():
    return FileResponse('c:/RSK-master/resourses/web/index.html')



def get_file():
    # Укажи путь к файлу Excel
    file_path = "c:/RSK-master/resourses/xl/raspisanie.xlsx"
    if not os.path.exists(file_path):
        raise FileNotFoundError("Файл с расписанием не найден")
    return file_path

@app.get("/group1/")
def group1(group: str, day_week: int):
    try:
        # Загрузка файла
        file = get_file()
        wb = load_workbook(file, data_only=True)
        sheet = wb.active

        # Установим начальный диапазон проверки
        max_rows = 250
        max_columns = 250
        column_group = None

        # Функция поиска группы в таблице
        def find_group_in_range(max_rows, max_columns):
            for i in range(1, max_rows + 1):  # Перебор строк
                for j in range(1, max_columns + 1):  # Перебор столбцов
                    if sheet.cell(row=i, column=j).value == group:
                        return j
            return None

        # Первичная попытка найти группу в ограниченном диапазоне
        column_group = find_group_in_range(max_rows, max_columns)

        # Если группа не найдена, увеличиваем диапазон до максимального размера таблицы
        if not column_group:
            max_rows = sheet.max_row
            max_columns = sheet.max_column
            column_group = find_group_in_range(max_rows, max_columns)

        # Если группа не найдена даже в расширенном диапазоне
        if not column_group:
            return JSONResponse(
                content={
                    "error": f"Группа '{group}' не найдена в файле Excel. Проверен диапазон {max_rows} строк и {max_columns} столбцов."
                },
                status_code=404,
            )

        # Рассчёт диапазона строк для указанного дня недели
        day_week = day_week * 13 - 7
        arr = []
        a = []

        # Сбор данных
        for i in range(day_week, day_week + 13):
            if sheet.cell(row=i, column=column_group + 1).value:
                a.append(sheet.cell(row=i, column=2).value)  # Время
                a.append(sheet.cell(row=i, column=3).value)  # Аудитория
                a.append(sheet.cell(row=i, column=column_group).value)  # Дисциплина
                a.append(sheet.cell(row=i, column=column_group + 1).value)  # Преподаватель
                a.append(sheet.cell(row=i, column=column_group + 2).value)  # Примечания

                arr.append(a.copy())  # Копируем данные строки
                a.clear()  # Очищаем временный массив

        return JSONResponse(content={"status": "success", "data": arr}, status_code=200)

    except FileNotFoundError as e:
        return JSONResponse(content={"error": str(e)}, status_code=500)

    except Exception as e:
        return JSONResponse(
            content={"error": "Внутренняя ошибка сервера", "details": str(e)},
            status_code=500,
        )



@app.get("/authorized")
async def read_auth():
    return FileResponse('c:/RSK-master/resourses/web/authorized.html')


@app.get("/teachers")
async def read_teachers():
    return FileResponse('resourses/web/teachers.html')

if __name__ == '__main__':
    uvicorn.run("maen:app", reload= True)