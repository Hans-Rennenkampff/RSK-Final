# main.py
from fastapi import FastAPI, HTTPException, Query, Form, Depends, Header
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles as SF
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from jose import JWTError, jwt
from openpyxl import load_workbook
import uvicorn
import os
from pathlib import Path 
from database import SessionLocal, User
from auth import hash_password, verify_password, create_access_token

SECRET_KEY = "secret"
ALGORITHM = "HS256"

app = FastAPI()
BASE_DIR = Path(__file__).resolve().parent
# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

@app.post("/register")
async def register(login: str = Form(...), password: str = Form(...), group: str = Form(...), db: Session = Depends(get_db)):
    if db.query(User).filter(User.username == login).first():
        raise HTTPException(status_code=400, detail="Пользователь уже существует")
    new_user = User(username=login, password_hash=hash_password(password), group=group)
    db.add(new_user)
    db.commit()
    return JSONResponse(content={"message": "Регистрация успешна!"})

@app.post("/login")
async def login(username: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == username).first()
    if user is None or not verify_password(password, user.password_hash):
        raise HTTPException(status_code=400, detail="Неверный логин или пароль")
    access_token = create_access_token(data={"sub": user.username})
    return JSONResponse(content={"token": access_token})

@app.get("/user-info")
async def get_user_info(Authorization: str = Header(...), db: Session = Depends(get_db)):
    token = Authorization.split(" ")[1]
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        user = db.query(User).filter(User.username == username).first()
        if user is None:
            raise HTTPException(status_code=404, detail="Пользователь не найден")
        return {"username": user.username, "group": user.group}
    except JWTError:
        raise HTTPException(status_code=401, detail="Ошибка токена")

@app.get("/authorized")
async def authorized_page(Authorization: str = Header(...)):
    token = Authorization.split(" ")[1]
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        return FileResponse(BASE_DIR / "resourses"/"web"/"authorized.html")
    except JWTError:
        raise HTTPException(status_code=401, detail="Ошибка токена")


@app.get("/group1")
def group1(group: str, day_week: int):
    try:
        file_path = BASE_DIR / "resourses"/"xl"/"raspisanie.xlsx"
        if not os.path.exists(file_path):
            raise FileNotFoundError("Файл с расписанием не найден")

        wb = load_workbook(file_path, data_only=True)
        sheet = wb.active
        max_rows = sheet.max_row
        max_columns = sheet.max_column

        def find_group():
            for i in range(1, max_rows + 1):
                for j in range(1, max_columns + 1):
                    if sheet.cell(row=i, column=j).value == group:
                        return j
            return None

        column_group = find_group()

        if not column_group:
            return JSONResponse(
                content={"error": f"Группа '{group}' не найдена."},
                status_code=404
            )

        day_week = day_week * 13 - 7
        result = []

        for i in range(day_week, day_week + 13):
            row_data = []
            if sheet.cell(row=i, column=column_group + 1).value:
                row_data.append(sheet.cell(row=i, column=2).value)  
                row_data.append(sheet.cell(row=i, column=3).value)  
                row_data.append(sheet.cell(row=i, column=column_group).value)  
                row_data.append(sheet.cell(row=i, column=column_group + 1).value)  
                row_data.append(sheet.cell(row=i, column=column_group + 2).value)  
                result.append(row_data)

        return JSONResponse(content={"status": "success", "data": result}, status_code=200)

    except FileNotFoundError as e:
        return JSONResponse(content={"error": str(e)}, status_code=500)

    except Exception as e:
        return JSONResponse(
            content={"error": "Внутренняя ошибка сервера", "details": str(e)},
            status_code=500
        )

@app.get("/")
async def read_index():
    return FileResponse(BASE_DIR / "resourses"/"web"/"index.html")
    

@app.get("/teachers")
async def read_teachers():
    return FileResponse(BASE_DIR / 'resourses'/'web'/'teachers.html')

app.mount("/static", SF(directory= BASE_DIR / "resourses"/"web"), name="static")



if __name__ == '__main__':
    uvicorn.run("main:app", reload= True)